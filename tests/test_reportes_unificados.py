"""Pruebas del contrato común de datos para los informes exportables."""

import datetime
import json
import re
import sqlite3

import pytest
from openpyxl import load_workbook

from services.reportes_unificados import pdf as pdf_reportes
from services.reportes_unificados.datos import (
    crear_datos_diarios,
    crear_datos_detalle,
    crear_datos_mensuales,
    crear_datos_objetivo,
)
from services.reportes_unificados.excel import exportar_excel
from services.reportes_unificados.pdf import exportar_pdf
from services.reportes_unificados import archivos as archivos_reportes
from services.reportes_unificados.modelos import (
    ColumnaReporte,
    MetadatosReporte,
    ReporteDatos,
)


def test_datos_mensuales_conservan_cifras_y_calculan_totales_una_vez():
    reporte = {
        "anio": 2026,
        "mes": 10,
        "cumplimiento_total": 50.0,
        "objetivos": [
            {
                "nombre": "Objetivo A",
                "dias_esperados": 4,
                "dias_con_pasada": 3,
                "dias_sin_pasada": 1,
                "cumplimiento_porcentaje": 75.0,
            },
            {
                "nombre": "Objetivo B",
                "dias_esperados": 2,
                "dias_con_pasada": 0,
                "dias_sin_pasada": 2,
                "cumplimiento_porcentaje": 0.0,
            },
        ],
    }

    datos = crear_datos_mensuales(reporte, {"turno": "Todos"})

    assert datos.filas[0]["cumplimiento"] == 0.75
    assert datos.filas[1]["dias_con_pasada"] == 0
    assert datos.totales == {
        "objetivos": 2,
        "dias_esperados": 6,
        "dias_con_pasada": 3,
        "dias_sin_pasada": 3,
        "cumplimiento_total": 0.5,
    }
    assert datos.metadatos.filtros_aplicados == {"turno": "Todos"}
    assert datos.metadatos.tipo_informe == "mensual"


def test_datos_diarios_preservan_filas_y_total_del_servicio():
    datos = crear_datos_diarios(
        {
            "fecha": "2026-10-31",
            "pasadas": [
                {
                    "hora": "23:50",
                    "turno": "nocturno",
                    "objetivo": "Objetivo Ñ",
                    "supervisor": "Muñoz",
                }
            ],
            "total": 1,
            "cumplimiento_porcentaje": 100.0,
        }
    )

    assert datos.filas[0]["fecha"] == datetime.date(2026, 10, 31)
    assert datos.filas[0]["hora"] == datetime.time(23, 50)
    assert datos.totales["pasadas"] == 1
    assert datos.filas[0]["objetivo"] == "Objetivo Ñ"


def test_datos_por_objetivo_usan_todos_los_totales_del_calculo():
    datos = crear_datos_objetivo(
        2026,
        10,
        {
            "nombre": "Objetivo de prueba",
            "dias": [
                {
                    "fecha": "31/10/2026",
                    "dia_semana": "Sáb",
                    "diurno": "A. Pérez",
                    "nocturno": "✘",
                    "estado": "Solo diurno",
                }
            ],
            "resumen": {
                "dias_esperados": 1,
                "dias_con_dia": 1,
                "dias_con_noche": 0,
                "dias_sin_control": 0,
                "porcentaje": 100.0,
            },
        },
    )

    assert datos.filas[0]["diurno"] == "A. Pérez"
    assert datos.totales["dias_con_noche"] == 0
    assert datos.metadatos.periodo == "2026-10"


def test_reporte_rechaza_ids_de_columna_duplicados():
    metadatos = MetadatosReporte(
        titulo="Prueba",
        tipo_informe="diario",
        periodo="2026-10-31",
        filtros_aplicados={},
        criterio_conteo_turnos="Criterio",
        generado_en=datetime.datetime(2026, 10, 31, 9, 0),
        usuario="Prueba",
        version_aplicacion="1.0",
    )
    columna = ColumnaReporte("fecha", "Fecha", "fecha", "dd/mm/aaaa", 14)

    with pytest.raises(ValueError, match="IDs de columna repetidos"):
        ReporteDatos(metadatos, (columna, columna), ())


def test_detalle_atribuye_madrugada_al_mes_de_turno(
    db_initialized, test_objetivo, test_supervisor
):
    conexion = sqlite3.connect(db_initialized)
    conexion.execute(
        """
        INSERT INTO pasadas
            (fecha, hora, turno, objetivo_id, supervisor_id, fecha_operativa)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            "2026-11-01",
            "02:00",
            "nocturno",
            test_objetivo["id"],
            test_supervisor["id"],
            "2026-10-31",
        ),
    )
    conexion.commit()
    conexion.close()

    datos = crear_datos_detalle(
        datetime.date(2026, 10, 31),
        datetime.date(2026, 10, 31),
    )

    assert len(datos.filas) == 1
    assert datos.filas[0]["fecha"] == datetime.date(2026, 10, 31)
    assert datos.totales == {
        "pasadas": 1,
        "pasadas_diurnas": 0,
        "pasadas_nocturnas": 1,
        "por_objetivo": {
            test_objetivo["id"]: {"diurno": 0, "nocturno": 1}
        },
    }


def test_detalle_aplica_filtros_de_objetivo_supervisor_y_turno(
    db_initialized, test_objetivo, test_supervisor
):
    conexion = sqlite3.connect(db_initialized)
    cursor = conexion.cursor()
    cursor.execute(
        "INSERT INTO objetivos (nombre, dias_semana) VALUES (?, ?)",
        ("Objetivo alternativo", "1,2,3,4,5"),
    )
    objetivo_alternativo = cursor.lastrowid
    cursor.execute("INSERT INTO supervisores (nombre) VALUES (?)", ("Supervisor B",))
    supervisor_alternativo = cursor.lastrowid
    cursor.executemany(
        """
        INSERT INTO pasadas (fecha, hora, turno, objetivo_id, supervisor_id, fecha_operativa)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        [
            (
                "2026-10-30",
                "10:00",
                "diurno",
                test_objetivo["id"],
                test_supervisor["id"],
                "2026-10-30",
            ),
            (
                "2026-10-30",
                "22:00",
                "nocturno",
                objetivo_alternativo,
                supervisor_alternativo,
                "2026-10-30",
            ),
        ],
    )
    conexion.commit()
    conexion.close()

    datos = crear_datos_detalle(
        datetime.date(2026, 10, 30),
        datetime.date(2026, 10, 30),
        {
            "objetivo_ids": [test_objetivo["id"]],
            "supervisor_ids": [test_supervisor["id"]],
            "turnos": ["diurno"],
        },
    )
    sin_coincidencias = crear_datos_detalle(
        datetime.date(2026, 10, 30),
        datetime.date(2026, 10, 30),
        {"supervisor_ids": [], "turnos": []},
    )

    assert len(datos.filas) == 1
    assert datos.filas[0]["objetivo"] == test_objetivo["nombre"]
    assert len(sin_coincidencias.filas) == 0


def test_resumen_mensual_usa_la_misma_seleccion_de_filtros(
    db_initialized, test_objetivo, test_supervisor
):
    from services.reportes_unificados.consultas import crear_datos_para_exportar

    with sqlite3.connect(db_initialized) as conexion:
        cursor = conexion.execute(
            "INSERT INTO objetivos (nombre, dias_semana) VALUES (?, ?)",
            ("Objetivo sin seleccionar", "1,2,3,4,5"),
        )
        otro_objetivo = cursor.lastrowid
        conexion.execute(
            "INSERT INTO supervisores (nombre) VALUES (?)",
            ("Supervisor excluido",),
        )
        otro_supervisor = conexion.execute(
            "SELECT id FROM supervisores WHERE nombre = ?",
            ("Supervisor excluido",),
        ).fetchone()[0]
        conexion.executemany(
            """
            INSERT INTO pasadas
                (fecha, fecha_operativa, hora, turno, objetivo_id, supervisor_id)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    "2026-10-28",
                    "2026-10-28",
                    "10:00",
                    "diurno",
                    test_objetivo["id"],
                    test_supervisor["id"],
                ),
                (
                    "2026-10-28",
                    "2026-10-28",
                    "22:00",
                    "nocturno",
                    test_objetivo["id"],
                    otro_supervisor,
                ),
                (
                    "2026-10-28",
                    "2026-10-28",
                    "11:00",
                    "diurno",
                    otro_objetivo,
                    test_supervisor["id"],
                ),
            ],
        )

    datos = crear_datos_para_exportar(
        {
            "tipo_informe": "mensual",
            "fecha_inicio": datetime.date(2026, 10, 1),
            "fecha_fin": datetime.date(2026, 10, 31),
            "objetivo_ids": [test_objetivo["id"]],
            "supervisor_ids": [test_supervisor["id"]],
            "turnos": ["diurno"],
            "equipo_id": None,
            "agrupacion": "supervisor",
        }
    )

    assert [fila["objetivo"] for fila in datos.filas] == [test_objetivo["nombre"]]
    assert datos.filas[0]["dias_con_pasada"] == 1
    assert datos.filas[0]["pasadas_diurnas"] == 1
    assert datos.filas[0]["pasadas_nocturnas"] == 0
    assert len(datos.filas_detalle) == 1
    assert datos.subtotales == (
        {
            "agrupacion": "supervisor",
            "valor": test_supervisor["nombre"],
            "etiqueta": test_supervisor["nombre"],
            "pasadas": 1,
        },
    )


def test_excel_usa_el_objeto_comun_respetando_orden_y_tipos(tmp_path):
    datos = crear_datos_mensuales(
        {
            "anio": 2026,
            "mes": 10,
            "cumplimiento_total": 75.0,
            "objetivos": [
                {
                    "id": 9,
                    "nombre": "Objetivo Ñ",
                    "dias_esperados": 4,
                    "dias_con_pasada": 3,
                    "dias_sin_pasada": 1,
                    "cumplimiento_porcentaje": 75.0,
                }
            ],
        },
        {"turno": ["diurno", "nocturno"]},
        filas_detalle=[
            {
                "fecha": datetime.date(2026, 10, 31),
                "hora": datetime.time(23, 50),
                "turno": "nocturno",
                "objetivo": "Objetivo Ñ",
                "supervisor": "Muñoz",
                "notas": 'Observación "larga" con tilde',
            }
        ],
        conteos_por_objetivo={9: {"diurno": 2, "nocturno": 3}},
    )
    ruta = tmp_path / "reporte.xlsx"

    exportar_excel(
        datos,
        ruta,
        columnas=[
            "objetivo",
            "pasadas_totales",
            "cumplimiento",
            "fecha",
            "hora",
            "notas",
        ],
        contenido="ambos",
        incluir_notas=True,
    )

    libro = load_workbook(ruta, data_only=True)
    resumen = libro["Resumen"]
    detalle = libro["Detalle"]
    assert [resumen.cell(5, col).value for col in range(1, 3)] == [
        "Objetivo",
        "Total de pasadas",
    ]
    assert resumen.cell(5, 3).value == "Cumplimiento"
    assert resumen.cell(6, 1).value == "Objetivo Ñ"
    assert resumen.cell(6, 2).value == 5
    assert resumen.cell(6, 2).data_type == "n"
    assert resumen.cell(6, 3).value == 0.75
    assert resumen.cell(6, 3).number_format == "0.0%"
    assert resumen.cell(7, 2).value == 5
    assert [detalle.cell(5, col).value for col in range(1, 5)] == [
        "Objetivo",
        "Fecha",
        "Hora",
        "Observaciones",
    ]
    assert detalle.cell(6, 1).value == "Objetivo Ñ"
    assert detalle.cell(6, 2).value.date() == datetime.date(2026, 10, 31)
    assert detalle.cell(6, 3).value == datetime.time(23, 50)
    assert detalle.cell(6, 4).value == 'Observación "larga" con tilde'
    assert detalle.cell(6, 2).number_format == "dd/mm/yyyy"
    assert detalle.cell(6, 3).number_format == "hh:mm"
    assert libro.sheetnames == ["Resumen", "Detalle", "Parámetros"]
    assert resumen.freeze_panes == "A6"
    assert resumen.auto_filter.ref == "A5:C6"
    assert resumen.print_title_rows == "$5:$5"
    assert resumen.page_setup.fitToWidth == 1
    assert resumen.column_dimensions["A"].width >= 34


def test_excel_y_pdf_conservan_los_mismos_datos_en_tablas_largas(
    tmp_path, monkeypatch
):
    reporte = {
        "anio": 2026,
        "mes": 10,
        "cumplimiento_total": 60.0,
        "objetivos": [
            {
                "id": identificador,
                "nombre": f"Objetivo Ñ {identificador} — " + ("texto largo " * 18),
                "dias_esperados": 10,
                "dias_con_pasada": 6,
                "dias_sin_pasada": 4,
                "cumplimiento_porcentaje": 60.0,
            }
            for identificador in range(1, 81)
        ],
    }
    datos = crear_datos_mensuales(reporte)
    ruta_pdf = tmp_path / "reporte.pdf"
    ruta_excel = tmp_path / "reporte.xlsx"
    valores_pdf = []
    texto_celda_original = pdf_reportes._texto_celda

    def registrar_texto(valor, columna):
        texto = texto_celda_original(valor, columna)
        valores_pdf.append(texto)
        return texto

    monkeypatch.setattr(pdf_reportes, "_texto_celda", registrar_texto)

    exportar_excel(datos, ruta_excel, contenido="resumen")
    exportar_pdf(datos, ruta_pdf, contenido="resumen")

    contenido_pdf = ruta_pdf.read_bytes()
    paginas = len(re.findall(rb"/Type\s*/Page\b", contenido_pdf))
    assert contenido_pdf.startswith(b"%PDF-")
    assert paginas >= 2
    assert len(datos.filas) == 80
    assert datos.filas[0]["objetivo"].startswith("Objetivo Ñ")
    assert any("Objetivo Ñ 1" in valor for valor in valores_pdf)
    assert "60.0%" in valores_pdf
    libro = load_workbook(ruta_excel, data_only=True)
    hoja = libro["Resumen"]
    assert hoja["C6"].value == 6
    assert hoja["C86"].value == 480
    libro.close()


def test_nombre_sugerido_sanitiza_caracteres_y_respeta_limite():
    ahora = datetime.datetime(2026, 10, 31, 14, 5)
    nombre = archivos_reportes.nombre_sugerido(
        "objetivo / Ñ", '2026-10: "especial"', "xlsx", ahora
    )

    assert nombre.endswith("_20261031-1405.xlsx")
    assert len(nombre) <= 120
    assert not any(caracter in nombre for caracter in '<>:"/\\|?*')


def test_guardar_reporte_no_sobrescribe_y_audita(tmp_path, monkeypatch):
    datos = crear_datos_mensuales(
        {
            "anio": 2026,
            "mes": 10,
            "cumplimiento_total": 100.0,
            "objetivos": [
                {
                    "nombre": "Objetivo",
                    "dias_esperados": 1,
                    "dias_con_pasada": 1,
                    "dias_sin_pasada": 0,
                    "cumplimiento_porcentaje": 100.0,
                }
            ],
        }
    )
    auditos = []
    monkeypatch.setattr(archivos_reportes, "tiene_permiso", lambda _: True)
    monkeypatch.setattr(
        archivos_reportes, "_registrar_exportacion", lambda *args: auditos.append(args)
    )
    ruta = tmp_path / "reporte.xlsx"
    archivos_reportes.guardar_reporte(
        datos, ruta, "excel", contenido="resumen"
    )

    assert ruta.exists()
    assert len(auditos) == 1
    with pytest.raises(FileExistsError, match="Ya existe"):
        archivos_reportes.guardar_reporte(
            datos, ruta, "excel", contenido="resumen"
        )
    assert len(auditos) == 1


def test_guardar_reporte_rechaza_usuario_sin_permiso(tmp_path, monkeypatch):
    monkeypatch.setattr(archivos_reportes, "tiene_permiso", lambda _: False)

    with pytest.raises(PermissionError, match="No tiene permiso"):
        archivos_reportes.guardar_reporte(
            None, tmp_path / "reporte.xlsx", "excel"
        )


def test_error_de_permiso_al_escribir_se_traduce_y_no_deja_archivo(
    tmp_path, monkeypatch
):
    datos = crear_datos_mensuales(
        {
            "anio": 2026,
            "mes": 10,
            "cumplimiento_total": 100.0,
            "objetivos": [
                {
                    "nombre": "Objetivo",
                    "dias_esperados": 1,
                    "dias_con_pasada": 1,
                    "dias_sin_pasada": 0,
                    "cumplimiento_porcentaje": 100.0,
                }
            ],
        }
    )
    monkeypatch.setattr(archivos_reportes, "tiene_permiso", lambda _: True)
    monkeypatch.setattr(
        archivos_reportes,
        "_registrar_exportacion",
        lambda *args: None,
    )
    monkeypatch.setattr(
        archivos_reportes.os,
        "open",
        lambda *args, **kwargs: (_ for _ in ()).throw(PermissionError("ocupado")),
    )
    ruta = tmp_path / "ocupado.xlsx"

    with pytest.raises(PermissionError, match="abierto en Excel"):
        archivos_reportes.guardar_reporte(
            datos, ruta, "excel", contenido="resumen"
        )
    assert not ruta.exists()


def test_registrar_exportacion_guarda_usuario_filtros_y_archivo(
    tmp_path, monkeypatch
):
    from database import db
    import services.sesion as sesion

    ruta_db = tmp_path / "auditoria.sqlite"
    with sqlite3.connect(ruta_db) as conexion:
        conexion.execute("CREATE TABLE usuarios (id INTEGER, username TEXT)")
        conexion.execute("INSERT INTO usuarios VALUES (27, 'usuario_prueba')")
        conexion.execute(
            """
            CREATE TABLE auditoria (
                id INTEGER PRIMARY KEY,
                fecha TEXT, hora TEXT, usuario_id INTEGER, tipo_operacion TEXT,
                tabla TEXT, detalles TEXT, estado TEXT
            )
            """
        )
    monkeypatch.setattr(db, "DB_PATH", str(ruta_db))
    monkeypatch.setattr(sesion, "get_usuario_id", lambda: 27)
    datos = crear_datos_mensuales(
        {
            "anio": 2026,
            "mes": 10,
            "cumplimiento_total": 100.0,
            "objetivos": [],
        },
        filtros={"supervisor_ids": [27]},
    )
    ruta = tmp_path / "reporte.xlsx"

    archivos_reportes._registrar_exportacion(
        datos, ruta, "excel", ["objetivo"], "resumen"
    )

    with sqlite3.connect(ruta_db) as conexion:
        fila = conexion.execute(
            "SELECT usuario_id, tipo_operacion, tabla, detalles, estado "
            "FROM auditoria"
        ).fetchone()
    detalles = json.loads(fila[3])
    assert fila[:3] == (27, "EXPORTACION", "reportes")
    assert detalles["filtros"] == {"supervisor_ids": [27]}
    assert detalles["archivo"] == str(ruta)
    assert fila[4] == "EXITOSO"


def test_calculo_por_objetivo_usa_el_servicio_compartido(
    db_initialized, test_objetivo, admin_user, test_supervisor
):
    from services.reportes_unificados.objetivo import calcular_reporte_objetivo

    with sqlite3.connect(db_initialized) as conexion:
        conexion.execute(
            """
            INSERT INTO pasadas
                (fecha, fecha_operativa, hora, turno, objetivo_id, supervisor_id)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                "2026-10-28",
                "2026-10-28",
                "10:00",
                "diurno",
                test_objetivo["id"],
                test_supervisor["id"],
            ),
        )

    reporte = calcular_reporte_objetivo(
        2026,
        10,
        test_objetivo["id"],
        ruta_db=db_initialized,
    )

    fila = next(
        dia for dia in reporte["dias"] if dia["fecha"] == "28/10/2026"
    )
    assert fila["estado"] == "Solo diurno"
    assert "Supervisor Test" in fila["diurno"]
