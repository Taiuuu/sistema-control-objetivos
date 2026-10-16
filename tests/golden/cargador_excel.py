"""Carga conteos provisionales por objetivo y turno desde un Excel mensual."""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from services.calculo_pasadas import (
    fecha_calendario_de_turno,
    fecha_de_turno,
    normalizar_turno,
)
from services.importador.normalizador import normalizar_hora
from services.importador.parser import parsear_nombre_hoja


def cargar_conteos_excel(ruta: Path, anio: int, mes: int) -> dict[str, dict[str, int]]:
    """Cuenta pasadas natural-únicas de hojas del mes; la hoja define el turno."""
    libro = load_workbook(ruta, data_only=True, read_only=True)
    claves: set[tuple[str, str, datetime, str]] = set()

    try:
        for hoja in libro.worksheets:
            fecha_hoja, turno_hoja, problema = parsear_nombre_hoja(hoja.title, anio)
            if problema or fecha_hoja is None or turno_hoja is None:
                continue
            if fecha_hoja.month != mes:
                continue

            filas = hoja.iter_rows(values_only=True)
            encabezado: tuple[int, list[int]] | None = None
            for numero_fila, fila in enumerate(filas, start=1):
                valores = list(fila)
                if encabezado is None and numero_fila <= 15:
                    inicios = [
                        columna
                        for columna in range(max(0, len(valores) - 5))
                        if [
                            str(valor or "").strip().upper()
                            for valor in valores[columna : columna + 6]
                        ]
                        == ["NO", "OBJETIVO", "TURNO", "MOVIL", "HORA", "SUPERVISOR"]
                    ]
                    if inicios:
                        encabezado = (numero_fila, inicios)
                    continue

                if encabezado is None:
                    continue
                fila_encabezado, inicios = encabezado
                if numero_fila <= fila_encabezado:
                    continue

                for columna in inicios:
                    if columna >= len(valores):
                        continue
                    numero = valores[columna]
                    objetivo = valores[columna + 1] if columna + 1 < len(valores) else None
                    hora_cruda = valores[columna + 4] if columna + 4 < len(valores) else None
                    supervisor = (
                        valores[columna + 5] if columna + 5 < len(valores) else None
                    )
                    if not isinstance(numero, (int, float)) or not objetivo:
                        continue
                    hora = normalizar_hora(hora_cruda).hora
                    if hora is None:
                        continue

                    turno = normalizar_turno(turno_hoja)
                    fecha_calendario = fecha_calendario_de_turno(
                        fecha_hoja, hora, turno
                    )
                    fecha_hora = datetime.combine(fecha_calendario, hora)
                    fecha_turno = fecha_de_turno(fecha_hora, turno)
                    if (fecha_turno.year, fecha_turno.month) != (anio, mes):
                        continue
                    objetivo_normalizado = re.sub(
                        r"\s+", " ", str(objetivo)
                    ).strip()
                    supervisor_normalizado = re.sub(
                        r"\s+", " ", str(supervisor or "")
                    ).strip()
                    claves.add(
                        (
                            objetivo_normalizado,
                            turno,
                            fecha_hora,
                            supervisor_normalizado,
                        )
                    )
    finally:
        libro.close()

    resultado: dict[str, dict[str, int]] = defaultdict(
        lambda: {"diurno": 0, "nocturno": 0}
    )
    for objetivo, turno, _, _ in claves:
        resultado[objetivo][turno] += 1
    return {nombre: resultado[nombre] for nombre in sorted(resultado)}
