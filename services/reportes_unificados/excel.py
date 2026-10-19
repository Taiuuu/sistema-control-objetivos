"""Generación profesional de libros XLSX desde el modelo común del informe."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from services.reportes_unificados.modelos import ColumnaReporte, ReporteDatos

VERDE_INSTITUCIONAL = "0E7F09"
GRAFITO_INSTITUCIONAL = "231F20"
COLOR_FILA_ALTERNADA = "F3F6F2"
COLOR_TOTAL = "DCEBD9"
COLOR_BORDE = "D7D2D3"
FILA_CABECERA = 5


def exportar_excel(
    datos: ReporteDatos,
    ruta: str | Path,
    columnas: Sequence[str] | None = None,
    contenido: str = "ambos",
    incluir_totales: bool = True,
    incluir_notas: bool = False,
) -> None:
    """Escribe un libro desde una única instantánea de datos ya calculados."""
    if contenido not in {"resumen", "detalle", "ambos"}:
        raise ValueError(f"Contenido de informe no reconocido: {contenido}")
    if columnas is not None:
        disponibles = {
            columna.id
            for columna in (*datos.columnas, *datos.columnas_detalle)
        }
        desconocidas = set(columnas) - disponibles
        if desconocidas:
            raise ValueError(
                "Columnas no disponibles en este informe: "
                + ", ".join(sorted(desconocidas))
            )
        if len(columnas) != len(set(columnas)):
            raise ValueError("La selección de columnas contiene IDs repetidos")

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    filas_resumen = (
        datos.filas if contenido in {"resumen", "ambos"} else ()
    )
    _escribir_hoja(
        resumen,
        datos,
        datos.columnas,
        filas_resumen,
        columnas,
        incluir_totales,
        incluir_notas,
        datos.totales,
        "Resumen no solicitado; el detalle está disponible en la hoja Detalle.",
    )

    if contenido in {"detalle", "ambos"}:
        detalle = libro.create_sheet("Detalle")
        catalogo = datos.columnas_detalle or datos.columnas
        filas = datos.filas_detalle or datos.filas
        _escribir_hoja(
            detalle,
            datos,
            catalogo,
            filas,
            columnas,
            incluir_totales,
            incluir_notas,
            datos.totales,
            "No hay filas de detalle para los filtros seleccionados.",
        )

    parametros = libro.create_sheet("Parámetros")
    _escribir_parametros(parametros, datos, contenido, incluir_totales, incluir_notas)
    libro.active = 0
    libro.save(ruta)


def _escribir_hoja(
    hoja: Worksheet,
    datos: ReporteDatos,
    catalogo: Sequence[ColumnaReporte],
    filas: Sequence[Mapping[str, Any]],
    ids_seleccionados: Sequence[str] | None,
    incluir_totales: bool,
    incluir_notas: bool,
    totales: Mapping[str, Any],
    mensaje_vacio: str,
) -> None:
    columnas = _seleccionar_columnas(catalogo, ids_seleccionados, incluir_notas)
    cantidad_columnas = max(1, len(columnas))
    ultima_columna = get_column_letter(cantidad_columnas)

    hoja.merge_cells(f"A1:{ultima_columna}1")
    titulo = hoja["A1"]
    titulo.value = "VESP Organizations"
    titulo.font = Font(bold=True, size=16, color=GRAFITO_INSTITUCIONAL)
    titulo.alignment = Alignment(horizontal="center", vertical="center")

    hoja.merge_cells(f"A2:{ultima_columna}2")
    hoja["A2"] = (
        f"{datos.metadatos.titulo} — {datos.metadatos.periodo}"
    )
    hoja["A2"].font = Font(bold=True, size=12, color=GRAFITO_INSTITUCIONAL)
    hoja["A2"].alignment = Alignment(horizontal="center", vertical="center")

    hoja.merge_cells(f"A3:{ultima_columna}3")
    hoja["A3"] = "Filtros: " + _texto_filtros(datos.metadatos.filtros_aplicados)
    hoja["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    hoja.row_dimensions[3].height = 30

    if not filas:
        hoja.merge_cells(f"A4:{ultima_columna}4")
        hoja["A4"] = mensaje_vacio
        _estilo_contenido(hoja, 4, cantidad_columnas)
        return

    if not columnas:
        raise ValueError("Debe seleccionarse al menos una columna para Excel")
    for columna_idx, columna in enumerate(columnas, start=1):
        celda = hoja.cell(FILA_CABECERA, columna_idx, columna.etiqueta)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE_INSTITUCIONAL)
        celda.alignment = Alignment(
            horizontal=_alineacion_excel(columna.alineacion),
            vertical="center",
            wrap_text=True,
        )

    for fila_idx, fila in enumerate(filas, start=FILA_CABECERA + 1):
        for columna_idx, columna in enumerate(columnas, start=1):
            celda = hoja.cell(
                fila_idx,
                columna_idx,
                _normalizar_valor(fila[columna.id], columna),
            )
            celda.alignment = Alignment(
                horizontal=_alineacion_excel(columna.alineacion),
                vertical="top",
                wrap_text=columna.tipo_dato in {"texto", "texto_largo"},
            )
            if columna.formato:
                celda.number_format = columna.formato
        if (fila_idx - FILA_CABECERA) % 2 == 0:
            for celda in hoja[fila_idx][:cantidad_columnas]:
                celda.fill = PatternFill("solid", fgColor=COLOR_FILA_ALTERNADA)

    if incluir_totales and datos.totales_incluidos:
        fila_total = FILA_CABECERA + len(filas) + 1
        for columna_idx, columna in enumerate(columnas, start=1):
            valor = _valor_total(columna, totales)
            celda = hoja.cell(
                fila_total,
                columna_idx,
                "Totales" if columna_idx == 1 and valor is None else valor,
            )
            celda.font = Font(bold=True, color=GRAFITO_INSTITUCIONAL)
            celda.fill = PatternFill("solid", fgColor=COLOR_TOTAL)
            celda.alignment = Alignment(
                horizontal=_alineacion_excel(columna.alineacion),
                vertical="top",
                wrap_text=True,
            )
            if columna.formato:
                celda.number_format = columna.formato
        if hoja.title == "Detalle" and "pasadas" in totales:
            fila_resumen_total = fila_total + 1
            hoja.merge_cells(
                f"A{fila_resumen_total}:{ultima_columna}{fila_resumen_total}"
            )
            hoja.cell(
                fila_resumen_total,
                1,
                _texto_totales_pasadas(totales),
            )
            hoja.cell(fila_resumen_total, 1).font = Font(
                bold=True, color=GRAFITO_INSTITUCIONAL
            )
            hoja.cell(fila_resumen_total, 1).fill = PatternFill(
                "solid", fgColor=COLOR_TOTAL
            )
            hoja.cell(fila_resumen_total, 1).alignment = Alignment(wrap_text=True)

    ultima_fila = FILA_CABECERA + len(filas)
    hoja.freeze_panes = f"A{FILA_CABECERA + 1}"
    hoja.auto_filter.ref = f"A{FILA_CABECERA}:{ultima_columna}{ultima_fila}"
    _ajustar_anchos(hoja, columnas, filas)
    _ajustar_altos(hoja, columnas, filas)
    _estilo_contenido(hoja, ultima_fila, cantidad_columnas)
    hoja.row_dimensions[1].height = 28
    hoja.row_dimensions[2].height = 24
    hoja.row_dimensions[FILA_CABECERA].height = 32

    ancho_estimado = sum(columna.ancho_sugerido for columna in columnas)
    hoja.page_setup.orientation = "landscape" if len(columnas) > 5 or ancho_estimado > 90 else "portrait"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = f"{FILA_CABECERA}:{FILA_CABECERA}"
    hoja.page_margins.left = 0.35
    hoja.page_margins.right = 0.35
    hoja.page_margins.top = 0.55
    hoja.page_margins.bottom = 0.55
    hoja.oddHeader.center.text = "&BVESP Organizations"
    hoja.oddFooter.left.text = f"Generado: {datos.metadatos.generado_en:%d/%m/%Y %H:%M}"
    hoja.oddFooter.right.text = "Página &[Page] de &[Pages]"


def _seleccionar_columnas(
    catalogo: Sequence[ColumnaReporte],
    ids: Sequence[str] | None,
    incluir_notas: bool,
) -> list[ColumnaReporte]:
    disponibles = {columna.id: columna for columna in catalogo}
    orden = list(ids) if ids is not None else [columna.id for columna in catalogo]
    if not incluir_notas:
        orden = [
            identificador
            for identificador in orden
            if identificador not in {"notas", "observaciones"}
        ]
    if len(orden) != len(set(orden)):
        raise ValueError("La selección de columnas contiene IDs repetidos")
    seleccionadas = [
        disponibles[identificador]
        for identificador in orden
        if identificador in disponibles
    ]
    return seleccionadas


def _normalizar_valor(valor: Any, columna: ColumnaReporte) -> Any:
    if columna.tipo_dato == "fecha_texto" and isinstance(valor, str):
        try:
            return datetime.strptime(valor, "%d/%m/%Y").date()
        except ValueError as error:
            raise ValueError(
                f"Fecha inválida para la columna {columna.etiqueta}: {valor!r}"
            ) from error
    if columna.tipo_dato == "entero" and valor is not None:
        return int(valor)
    if columna.tipo_dato in {"decimal", "porcentaje"} and valor is not None:
        return float(valor)
    if columna.tipo_dato == "fecha" and valor is not None and not isinstance(valor, date):
        raise TypeError(f"La columna {columna.etiqueta} requiere una fecha")
    if columna.tipo_dato == "hora" and valor is not None and not isinstance(valor, time):
        raise TypeError(f"La columna {columna.etiqueta} requiere una hora")
    return valor


def _valor_total(columna: ColumnaReporte, totales: Mapping[str, Any]) -> Any:
    aliases = {
        "cumplimiento": "cumplimiento_total",
        "pasadas_totales": "pasadas",
        "dias_con_pasada": "dias_con_pasada",
        "dias_sin_pasada": "dias_sin_pasada",
        "pasadas_diurnas": "pasadas_diurnas",
        "pasadas_nocturnas": "pasadas_nocturnas",
    }
    return totales.get(aliases.get(columna.id, columna.id))


def _ajustar_anchos(
    hoja: Worksheet,
    columnas: Sequence[ColumnaReporte],
    filas: Sequence[Mapping[str, Any]],
) -> None:
    for indice, columna in enumerate(columnas, start=1):
        largo_maximo = max(
            [len(columna.etiqueta)]
            + [
                len(str(fila.get(columna.id, "")))
                for fila in filas
                if fila.get(columna.id) is not None
            ]
        )
        ancho = max(columna.ancho_sugerido, min(largo_maximo + 2, 60))
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def _ajustar_altos(
    hoja: Worksheet,
    columnas: Sequence[ColumnaReporte],
    filas: Sequence[Mapping[str, Any]],
) -> None:
    for fila_idx, fila in enumerate(filas, start=FILA_CABECERA + 1):
        lineas_necesarias = 1
        for indice, columna in enumerate(columnas, start=1):
            if columna.tipo_dato not in {"texto", "texto_largo"}:
                continue
            texto = str(fila.get(columna.id) or "")
            ancho = max(1, int(hoja.column_dimensions[get_column_letter(indice)].width))
            lineas_necesarias = max(
                lineas_necesarias,
                sum(max(1, (len(linea) + ancho - 1) // ancho) for linea in texto.splitlines()),
            )
        hoja.row_dimensions[fila_idx].height = min(15 * lineas_necesarias, 400)


def _estilo_contenido(hoja: Worksheet, ultima_fila: int, columnas: int) -> None:
    borde = Border(
        bottom=Side(style="hair", color=COLOR_BORDE),
        right=Side(style="hair", color=COLOR_BORDE),
    )
    for fila in hoja.iter_rows(min_row=FILA_CABECERA, max_row=ultima_fila, max_col=columnas):
        for celda in fila:
            celda.border = borde


def _alineacion_excel(alineacion: str) -> str:
    return {"centro": "center", "derecha": "right"}.get(alineacion, "left")


def _texto_filtros(filtros: Mapping[str, Any]) -> str:
    if not filtros:
        return "Sin filtros adicionales"
    return " | ".join(f"{clave}: {valor}" for clave, valor in filtros.items())


def _texto_totales_pasadas(totales: Mapping[str, Any]) -> str:
    partes = [f"Total de pasadas: {totales['pasadas']}"]
    if "pasadas_diurnas" in totales:
        partes.append(f"Día: {totales['pasadas_diurnas']}")
    if "pasadas_nocturnas" in totales:
        partes.append(f"Noche: {totales['pasadas_nocturnas']}")
    return " | ".join(partes)


def _escribir_parametros(
    hoja: Worksheet,
    datos: ReporteDatos,
    contenido: str,
    incluir_totales: bool,
    incluir_notas: bool,
) -> None:
    hoja.append(["Parámetro", "Valor"])
    valores = [
        ("Título", datos.metadatos.titulo),
        ("Tipo de informe", datos.metadatos.tipo_informe),
        ("Período", datos.metadatos.periodo),
        ("Filtros aplicados", _texto_filtros(datos.metadatos.filtros_aplicados)),
        ("Criterio de conteo de turnos", datos.metadatos.criterio_conteo_turnos),
        ("Fecha y hora de generación", datos.metadatos.generado_en),
        ("Usuario", datos.metadatos.usuario),
        ("Versión de la aplicación", datos.metadatos.version_aplicacion),
        ("Contenido", contenido),
        ("Totales incluidos", "Sí" if incluir_totales else "No"),
        ("Notas incluidas", "Sí" if incluir_notas else "No"),
    ]
    for parametro, valor in valores:
        hoja.append([parametro, valor])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE_INSTITUCIONAL)
    hoja.column_dimensions["A"].width = 32
    hoja.column_dimensions["B"].width = 70
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    hoja.freeze_panes = "A2"
    hoja.page_setup.orientation = "portrait"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.oddFooter.left.text = f"Generado: {datos.metadatos.generado_en:%d/%m/%Y %H:%M}"
    hoja.oddFooter.right.text = "Página &[Page] de &[Pages]"
